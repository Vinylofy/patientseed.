"""Derive one public work item per shop row from the supplied workbook."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from urllib.parse import urlsplit

FIRST_SHOP_ROW = 7
LAST_SHOP_ROW = 3309
SHARED_HOSTS = {"facebook.com", "m.facebook.com", "instagram.com", "x.com", "twitter.com",
                "linktr.ee", "a.co", "amazon.com", "sites.google.com", "ebay.com", "etsy.com"}


def source_domain(source: object) -> tuple[str, str]:
    if not isinstance(source, str) or not source.strip():
        return "", "MISSING_SOURCE"
    try:
        parsed = urlsplit(source.strip())
        domain = (parsed.hostname or "").lower().removeprefix("www.")
        if (parsed.scheme not in {"http", "https"} or parsed.username or parsed.password
                or parsed.port not in (None, 80, 443) or not re.fullmatch(r"[a-z0-9.-]+", domain)
                or "." not in domain or len(domain) > 253):
            return "", "INVALID_SOURCE"
    except ValueError:
        return "", "INVALID_SOURCE"
    if domain in SHARED_HOSTS or domain.endswith((".facebook.com", ".instagram.com", ".amazon.com")):
        return domain, "SHARED_PLATFORM"
    return domain, "READY"


def build_rows(workbook: Path, bootstrap: list[str]) -> list[dict]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required only to regenerate the queue") from exc
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True, keep_links=False)
    try:
        if book.sheetnames != ["Unieke shops"]:
            raise ValueError("unexpected workbook sheets")
        sheet = book.active
        rows = []
        for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
            if number < FIRST_SHOP_ROW or number > LAST_SHOP_ROW:
                continue
            if len(values) < 5 or not isinstance(values[0], str) or not values[0].strip():
                raise ValueError(f"missing shop at workbook row {number}")
            domain, status = source_domain(values[4])
            rows.append({"row": number, "name": values[0].strip(), "domain": domain, "source_status": status})
    finally:
        book.close()
    if len(rows) != 3303:
        raise ValueError("shop count changed; review the workbook section before regenerating")
    by_domain = {}
    for item in rows:
        if item["source_status"] == "READY":
            by_domain.setdefault(item["domain"], item)
    if len(bootstrap) != len(set(bootstrap)) or any(domain not in by_domain for domain in bootstrap):
        raise ValueError("bootstrap domain missing or duplicated in workbook")
    selected = [by_domain[domain] for domain in bootstrap]
    selected_rows = {item["row"] for item in selected}
    ordered = selected + [item for item in rows if item["row"] not in selected_rows]
    seen: set[str] = set()
    for item in ordered:
        if item["source_status"] == "READY":
            if item["domain"] in seen:
                item["source_status"] = "DUPLICATE_DOMAIN"
            else:
                seen.add(item["domain"])
    return ordered


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--bootstrap", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    bootstrap = json.loads(args.bootstrap.read_text(encoding="utf-8"))
    rows = build_rows(args.workbook, bootstrap)
    args.output.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(rows)} shop rows to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
